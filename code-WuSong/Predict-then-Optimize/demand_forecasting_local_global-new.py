from __future__ import annotations

import json
import math
import time
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import HistGradientBoostingRegressor
from sklearn.linear_model import ElasticNet, LinearRegression, Ridge
from sklearn.metrics import mean_absolute_error, mean_squared_error
from sklearn.neural_network import MLPRegressor
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler

try:
    from catboost import CatBoostRegressor

    CATBOOST_AVAILABLE = True
except ImportError:
    CATBOOST_AVAILABLE = False

try:
    from xgboost import XGBRegressor

    XGBOOST_AVAILABLE = True
except ImportError:
    XGBOOST_AVAILABLE = False


# =========================
# Paths and experiment setup
# =========================
ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "upload" / "data-new.xlsx"
OUTPUT_DIR = ROOT / "outputs" / "demand_forecasting_actual_month_days"
VALIDATION_XLSX = OUTPUT_DIR / "模型验证结果.xlsx"
TEST_XLSX = OUTPUT_DIR / "模型测试结果.xlsx"

TRAIN_END = pd.Timestamp("2024-12-01")
VALIDATION_START = pd.Timestamp("2025-01-01")
VALIDATION_END = pd.Timestamp("2025-06-01")
TEST_START = pd.Timestamp("2025-07-01")
TEST_END = pd.Timestamp("2025-12-01")

RANDOM_STATE = 2026
RIDGE_ALPHAS = [0.01, 0.1, 1.0, 10.0, 100.0]
ELASTIC_NET_GRID = [
    {"alpha": alpha, "l1_ratio": l1_ratio}
    for alpha in [0.001, 0.01, 0.1, 1.0]
    for l1_ratio in [0.1, 0.5, 0.9]
]
CATBOOST_GRID = [
    {
        "depth": depth,
        "learning_rate": learning_rate,
        "iterations": iterations,
        "l2_leaf_reg": l2_leaf_reg,
    }
    for depth in [3, 4]
    for learning_rate in [0.03, 0.05]
    for iterations in [200, 300]
    for l2_leaf_reg in [5, 10]
]
XGBOOST_GRID = [
    {
        "max_depth": max_depth,
        "learning_rate": learning_rate,
        "n_estimators": n_estimators,
        "reg_lambda": reg_lambda,
    }
    for max_depth in [3, 4]
    for learning_rate in [0.03, 0.05]
    for n_estimators in [200, 300]
    for reg_lambda in [5, 10]
]
MLP_GRID = [
    {
        "hidden_layer_sizes": hidden_layer_sizes,
        "activation": activation,
        "alpha": alpha,
        "learning_rate_init": learning_rate_init,
    }
    for hidden_layer_sizes in [(16,), (32,), (32, 16)]
    for activation in ["relu", "tanh"]
    for alpha in [0.0001, 0.001]
    for learning_rate_init in [0.001, 0.005]
]

NUMERIC_FEATURES = [
    "lag_3",
    "lag_2",
    "lag_1",
    "month_sin",
    "month_cos",
    "spring_festival",
    "may_day",
    "national_day",
]
GLOBAL_FEATURES = NUMERIC_FEATURES + ["department"]

# 春节所在月份；五一和国庆分别固定在5月和10月。
SPRING_FESTIVAL_MONTH = {2023: 1, 2024: 2, 2025: 1, 2026: 2}


TRAINING_LOGS: list[dict] = []


def log(message: str) -> None:
    print(f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {message}", flush=True)


def load_excel_data(path: Path) -> pd.DataFrame:
    """Read wide monthly demand data and return department-month observations.

    Required columns: 科室, 床位数.
    Every column whose header can be parsed as YYYY-MM is treated as a monthly
    demand column. Monthly capacity is calculated as staffed beds multiplied
    by the actual number of calendar days in that month. New months can be
    appended to the input workbook without
    changing this function.
    """
    wide = pd.read_excel(path, sheet_name=0)
    wide.columns = [str(column).strip() for column in wide.columns]
    required = ["科室", "床位数"]
    missing = [column for column in required if column not in wide.columns]
    if missing:
        raise ValueError(f"{path.name} is missing required columns: {missing}")

    month_columns = []
    for column in wide.columns:
        try:
            parsed = pd.to_datetime(column, format="%Y-%m")
            month_columns.append((column, parsed))
        except (TypeError, ValueError):
            continue
    month_columns.sort(key=lambda item: item[1])
    month_names = [column for column, _ in month_columns]
    if len(month_names) < 4:
        raise ValueError("At least four YYYY-MM demand columns are required.")

    wide = wide[required + month_names].copy()
    wide["科室"] = wide["科室"].astype(str).str.strip()
    wide["床位数"] = pd.to_numeric(wide["床位数"], errors="raise")
    if (wide["床位数"] <= 0).any():
        raise ValueError("Staffed beds must be positive for every department.")

    long = wide.melt(
        id_vars=required,
        value_vars=month_names,
        var_name="date",
        value_name="demand",
    )
    long["date"] = pd.to_datetime(long["date"], format="%Y-%m")
    long["demand"] = pd.to_numeric(long["demand"], errors="coerce")
    if long["demand"].isna().any():
        bad = long.loc[long["demand"].isna(), ["科室", "date"]]
        raise ValueError(f"Missing or nonnumeric demand values found:\n{bad.to_string(index=False)}")
    data = long.rename(
        columns={"科室": "department", "床位数": "beds"}
    )[["department", "date", "beds", "demand"]]
    data["days_in_month"] = data["date"].dt.days_in_month.astype(int)
    data["capacity"] = data["beds"] * data["days_in_month"]
    data["load"] = data["demand"] / data["capacity"]
    data = data.sort_values(["department", "date"]).reset_index(drop=True)

    department_count = data["department"].nunique()
    month_count = data["date"].nunique()
    expected_rows = department_count * month_count
    if len(data) != expected_rows:
        raise ValueError(
            f"The table is not rectangular: expected {expected_rows} rows from "
            f"{department_count} departments x {month_count} months, found {len(data)}."
        )
    if not np.allclose(data["capacity"], data["beds"] * data["days_in_month"]):
        raise ValueError("Monthly-capacity consistency check failed.")
    log(
        f"Excel data loaded: sheet=Sheet1, departments={department_count}, "
        f"months={month_count}; capacity=beds*actual calendar days"
    )
    return data


def add_features(data: pd.DataFrame) -> pd.DataFrame:
    frames = []
    for department, group in data.groupby("department", sort=False):
        group = group.sort_values("date").copy()
        group["lag_1"] = group["load"].shift(1)
        group["lag_2"] = group["load"].shift(2)
        group["lag_3"] = group["load"].shift(3)
        frames.append(group)

    featured = pd.concat(frames, ignore_index=True).dropna(subset=["lag_1", "lag_2", "lag_3"])
    month = featured["date"].dt.month
    featured["month_sin"] = np.sin(2 * np.pi * month / 12)
    featured["month_cos"] = np.cos(2 * np.pi * month / 12)
    featured["spring_festival"] = [
        int(SPRING_FESTIVAL_MONTH.get(date.year) == date.month) for date in featured["date"]
    ]
    featured["may_day"] = (month == 5).astype(int)
    featured["national_day"] = (month == 10).astype(int)
    return featured.reset_index(drop=True)


def split_data(data: pd.DataFrame):
    train = data[data["date"] <= TRAIN_END].copy()
    validation = data[(data["date"] >= VALIDATION_START) & (data["date"] <= VALIDATION_END)].copy()
    test = data[(data["date"] >= TEST_START) & (data["date"] <= TEST_END)].copy()
    return train, validation, test


def metric_dict(actual: np.ndarray, predicted: np.ndarray) -> dict:
    actual = np.asarray(actual, dtype=float)
    predicted = np.asarray(predicted, dtype=float)
    error = predicted - actual
    denominator = float(np.sum(np.abs(actual)))
    return {
        "MAE": float(mean_absolute_error(actual, predicted)),
        "RMSE": float(math.sqrt(mean_squared_error(actual, predicted))),
        "WAPE": float(np.sum(np.abs(error)) / denominator) if denominator else np.nan,
        "Bias": float(np.sum(error) / denominator) if denominator else np.nan,
    }


def to_demand_prediction(load_prediction: np.ndarray, frame: pd.DataFrame) -> np.ndarray:
    return np.maximum(0.0, np.asarray(load_prediction, dtype=float)) * frame["capacity"].to_numpy()


def add_training_log(
    stage: str,
    structure: str,
    model: str,
    params: dict,
    elapsed: float,
    train_metrics: dict,
    validation_metrics: dict | None,
) -> None:
    row = {
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "stage": stage,
        "structure": structure,
        "model": model,
        "parameters": json.dumps(params, ensure_ascii=False, sort_keys=True),
        "elapsed_seconds": elapsed,
        "train_MAE": train_metrics["MAE"],
        "train_RMSE": train_metrics["RMSE"],
        "train_WAPE": train_metrics["WAPE"],
        "train_Bias": train_metrics["Bias"],
    }
    if validation_metrics is not None:
        row.update(
            {
                "validation_MAE": validation_metrics["MAE"],
                "validation_RMSE": validation_metrics["RMSE"],
                "validation_WAPE": validation_metrics["WAPE"],
                "validation_Bias": validation_metrics["Bias"],
            }
        )
    TRAINING_LOGS.append(row)


def make_local_estimator(model_name: str, params: dict):
    if model_name == "LinearRegression":
        estimator = LinearRegression()
    elif model_name == "Ridge":
        estimator = Ridge(alpha=params["alpha"])
    elif model_name == "ElasticNet":
        estimator = ElasticNet(
            alpha=params["alpha"],
            l1_ratio=params["l1_ratio"],
            max_iter=50000,
            random_state=RANDOM_STATE,
        )
    elif model_name == "DNN":
        estimator = MLPRegressor(
            hidden_layer_sizes=tuple(params["hidden_layer_sizes"]),
            activation=params["activation"],
            solver="adam",
            alpha=params["alpha"],
            learning_rate_init=params["learning_rate_init"],
            max_iter=3000,
            early_stopping=False,
            random_state=RANDOM_STATE,
        )
    else:
        raise ValueError(model_name)
    return Pipeline([("scale", StandardScaler()), ("model", estimator)])


def make_global_linear_estimator(model_name: str, params: dict):
    if model_name == "LinearRegression":
        estimator = LinearRegression()
    elif model_name == "Ridge":
        estimator = Ridge(alpha=params["alpha"])
    elif model_name == "ElasticNet":
        estimator = ElasticNet(
            alpha=params["alpha"],
            l1_ratio=params["l1_ratio"],
            max_iter=50000,
            random_state=RANDOM_STATE,
        )
    elif model_name == "DNN":
        estimator = MLPRegressor(
            hidden_layer_sizes=tuple(params["hidden_layer_sizes"]),
            activation=params["activation"],
            solver="adam",
            alpha=params["alpha"],
            learning_rate_init=params["learning_rate_init"],
            max_iter=3000,
            early_stopping=True,
            validation_fraction=0.15,
            n_iter_no_change=50,
            random_state=RANDOM_STATE,
        )
    else:
        raise ValueError(model_name)

    preprocess = ColumnTransformer(
        [
            ("numeric", StandardScaler(), NUMERIC_FEATURES),
            ("department", OneHotEncoder(handle_unknown="ignore", drop="first"), ["department"]),
        ],
        remainder="drop",
    )
    return Pipeline([("preprocess", preprocess), ("model", estimator)])


def make_tree_estimator(params: dict):
    if CATBOOST_AVAILABLE:
        return CatBoostRegressor(
            depth=params["depth"],
            learning_rate=params["learning_rate"],
            iterations=params["iterations"],
            l2_leaf_reg=params["l2_leaf_reg"],
            loss_function="RMSE",
            random_seed=RANDOM_STATE,
            verbose=False,
            allow_writing_files=False,
        )

    # Current environment fallback. The result is explicitly labelled and never presented as CatBoost.
    return HistGradientBoostingRegressor(
        max_iter=params["iterations"],
        learning_rate=params["learning_rate"],
        max_depth=params["depth"],
        l2_regularization=params["l2_leaf_reg"],
        min_samples_leaf=8,
        loss="squared_error",
        random_state=RANDOM_STATE,
    )


def make_xgboost_estimator(params: dict):
    if not XGBOOST_AVAILABLE:
        raise ImportError(
            "XGBoost is not installed. Install it with `python -m pip install xgboost` "
            "before running the XGBoost experiment."
        )
    return XGBRegressor(
        objective="reg:squarederror",
        max_depth=params["max_depth"],
        learning_rate=params["learning_rate"],
        n_estimators=params["n_estimators"],
        reg_lambda=params["reg_lambda"],
        reg_alpha=0.0,
        min_child_weight=8,
        subsample=0.9,
        colsample_bytree=0.9,
        tree_method="hist",
        n_jobs=-1,
        random_state=RANDOM_STATE,
        verbosity=0,
    )


def fit_local_predict(model_name: str, params: dict, fit_frame: pd.DataFrame, predict_frame: pd.DataFrame):
    fit_predictions = pd.Series(index=fit_frame.index, dtype=float)
    future_predictions = pd.Series(index=predict_frame.index, dtype=float)
    fitted_models = {}

    for department in sorted(fit_frame["department"].unique()):
        fit_dept = fit_frame[fit_frame["department"] == department]
        pred_dept = predict_frame[predict_frame["department"] == department]
        estimator = make_local_estimator(model_name, params)
        estimator.fit(fit_dept[NUMERIC_FEATURES], fit_dept["load"])
        fit_predictions.loc[fit_dept.index] = estimator.predict(fit_dept[NUMERIC_FEATURES])
        future_predictions.loc[pred_dept.index] = estimator.predict(pred_dept[NUMERIC_FEATURES])
        fitted_models[department] = estimator

    return fit_predictions.sort_index().to_numpy(), future_predictions.sort_index().to_numpy(), fitted_models


def fit_global_predict(model_name: str, params: dict, fit_frame: pd.DataFrame, predict_frame: pd.DataFrame):
    if model_name in {"LinearRegression", "Ridge", "ElasticNet", "DNN"}:
        estimator = make_global_linear_estimator(model_name, params)
        estimator.fit(fit_frame[GLOBAL_FEATURES], fit_frame["load"])
        return (
            estimator.predict(fit_frame[GLOBAL_FEATURES]),
            estimator.predict(predict_frame[GLOBAL_FEATURES]),
            estimator,
        )

    if model_name == "Tree":
        estimator = make_tree_estimator(params)
        if CATBOOST_AVAILABLE:
            x_fit = fit_frame[GLOBAL_FEATURES].copy()
            x_pred = predict_frame[GLOBAL_FEATURES].copy()
            estimator.fit(x_fit, fit_frame["load"], cat_features=["department"])
        else:
            # One-hot encode department for the sklearn fallback.
            preprocess = ColumnTransformer(
                [
                    ("numeric", "passthrough", NUMERIC_FEATURES),
                    (
                        "department",
                        OneHotEncoder(handle_unknown="ignore", sparse_output=False),
                        ["department"],
                    ),
                ]
            )
            estimator = Pipeline([("preprocess", preprocess), ("model", estimator)])
            estimator.fit(fit_frame[GLOBAL_FEATURES], fit_frame["load"])
        return (
            estimator.predict(fit_frame[GLOBAL_FEATURES]),
            estimator.predict(predict_frame[GLOBAL_FEATURES]),
            estimator,
        )
    if model_name == "XGBoost":
        preprocess = ColumnTransformer(
            [
                ("numeric", "passthrough", NUMERIC_FEATURES),
                ("department", OneHotEncoder(handle_unknown="ignore"), ["department"]),
            ]
        )
        estimator = Pipeline(
            [("preprocess", preprocess), ("model", make_xgboost_estimator(params))]
        )
        estimator.fit(fit_frame[GLOBAL_FEATURES], fit_frame["load"])
        return (
            estimator.predict(fit_frame[GLOBAL_FEATURES]),
            estimator.predict(predict_frame[GLOBAL_FEATURES]),
            estimator,
        )
    raise ValueError(model_name)


def prediction_frame(source: pd.DataFrame, structure: str, model: str, predicted_demand: np.ndarray) -> pd.DataFrame:
    result = source[
        ["date", "department", "beds", "days_in_month", "capacity", "demand", "load"]
    ].copy()
    result.insert(0, "structure", structure)
    result.insert(1, "model", model)
    result["predicted_demand"] = predicted_demand
    result["error"] = result["predicted_demand"] - result["demand"]
    result["absolute_error"] = result["error"].abs()
    result["predicted_load"] = result["predicted_demand"] / result["capacity"]
    return result


def department_metrics(predictions: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (structure, model, department), group in predictions.groupby(["structure", "model", "department"]):
        metrics = metric_dict(group["demand"], group["predicted_demand"])
        rows.append({"structure": structure, "model": model, "department": department, **metrics})
    return pd.DataFrame(rows).sort_values(["structure", "model", "department"])


def summary_metrics(predictions: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (structure, model), group in predictions.groupby(["structure", "model"]):
        metrics = metric_dict(group["demand"], group["predicted_demand"])
        rows.append({"structure": structure, "model": model, **metrics})
    return pd.DataFrame(rows).sort_values(["WAPE", "RMSE"])


def evaluate_baselines(validation: pd.DataFrame):
    outputs = []
    recent = validation["lag_1"].to_numpy() * validation["capacity"].to_numpy()
    moving_average = validation[["lag_1", "lag_2", "lag_3"]].mean(axis=1).to_numpy() * validation["capacity"].to_numpy()
    outputs.append(prediction_frame(validation, "Baseline", "RecentMonth", recent))
    outputs.append(prediction_frame(validation, "Baseline", "MovingAverage3", moving_average))
    return outputs


def tune_and_validate(train: pd.DataFrame, validation: pd.DataFrame):
    validation_predictions = evaluate_baselines(validation)
    selected = []

    model_grids = [
        ("LinearRegression", [{}]),
        ("Ridge", [{"alpha": alpha} for alpha in RIDGE_ALPHAS]),
        ("ElasticNet", ELASTIC_NET_GRID),
        ("DNN", MLP_GRID),
    ]

    for structure in ["Local", "Global"]:
        for model_name, grid in model_grids:
            best = None
            for params in grid:
                started = time.perf_counter()
                if structure == "Local":
                    train_load, val_load, _ = fit_local_predict(model_name, params, train, validation)
                else:
                    train_load, val_load, _ = fit_global_predict(model_name, params, train, validation)
                elapsed = time.perf_counter() - started
                train_pred = to_demand_prediction(train_load, train)
                val_pred = to_demand_prediction(val_load, validation)
                train_metrics = metric_dict(train["demand"], train_pred)
                val_metrics = metric_dict(validation["demand"], val_pred)
                add_training_log("validation_tuning", structure, model_name, params, elapsed, train_metrics, val_metrics)
                log(
                    f"{structure}-{model_name} params={params} | "
                    f"time={elapsed:.3f}s | train WAPE={train_metrics['WAPE']:.4f} | "
                    f"validation WAPE={val_metrics['WAPE']:.4f}"
                )
                candidate = (val_metrics["WAPE"], params, val_pred, val_metrics)
                if best is None or candidate[0] < best[0]:
                    best = candidate

            display_name = model_name
            validation_predictions.append(prediction_frame(validation, structure, display_name, best[2]))
            selected.append(
                {
                    "structure": structure,
                    "model": model_name,
                    "display_model": display_name,
                    "parameters": json.dumps(best[1], ensure_ascii=False, sort_keys=True),
                    "validation_WAPE": best[3]["WAPE"],
                }
            )

    tree_label = "CatBoost" if CATBOOST_AVAILABLE else "ShallowGBDT_Fallback"
    best = None
    for params in CATBOOST_GRID:
        started = time.perf_counter()
        train_load, val_load, _ = fit_global_predict("Tree", params, train, validation)
        elapsed = time.perf_counter() - started
        train_pred = to_demand_prediction(train_load, train)
        val_pred = to_demand_prediction(val_load, validation)
        train_metrics = metric_dict(train["demand"], train_pred)
        val_metrics = metric_dict(validation["demand"], val_pred)
        add_training_log("validation_tuning", "Global", tree_label, params, elapsed, train_metrics, val_metrics)
        log(
            f"Global-{tree_label} params={params} | time={elapsed:.3f}s | "
            f"train WAPE={train_metrics['WAPE']:.4f} | validation WAPE={val_metrics['WAPE']:.4f}"
        )
        candidate = (val_metrics["WAPE"], params, val_pred, val_metrics)
        if best is None or candidate[0] < best[0]:
            best = candidate

    validation_predictions.append(prediction_frame(validation, "Global", tree_label, best[2]))
    selected.append(
        {
            "structure": "Global",
            "model": "Tree",
            "display_model": tree_label,
            "parameters": json.dumps(best[1], ensure_ascii=False, sort_keys=True),
            "validation_WAPE": best[3]["WAPE"],
        }
    )

    if XGBOOST_AVAILABLE:
        best = None
        for params in XGBOOST_GRID:
            started = time.perf_counter()
            train_load, val_load, _ = fit_global_predict("XGBoost", params, train, validation)
            elapsed = time.perf_counter() - started
            train_pred = to_demand_prediction(train_load, train)
            val_pred = to_demand_prediction(val_load, validation)
            train_metrics = metric_dict(train["demand"], train_pred)
            val_metrics = metric_dict(validation["demand"], val_pred)
            add_training_log(
                "validation_tuning", "Global", "XGBoost", params, elapsed,
                train_metrics, val_metrics
            )
            log(
                f"Global-XGBoost params={params} | time={elapsed:.3f}s | "
                f"train WAPE={train_metrics['WAPE']:.4f} | "
                f"validation WAPE={val_metrics['WAPE']:.4f}"
            )
            candidate = (val_metrics["WAPE"], params, val_pred, val_metrics)
            if best is None or candidate[0] < best[0]:
                best = candidate

        validation_predictions.append(
            prediction_frame(validation, "Global", "XGBoost", best[2])
        )
        selected.append(
            {
                "structure": "Global",
                "model": "XGBoost",
                "display_model": "XGBoost",
                "parameters": json.dumps(best[1], ensure_ascii=False, sort_keys=True),
                "validation_WAPE": best[3]["WAPE"],
            }
        )
    else:
        log(
            "XGBoost skipped: package is not installed. "
            "Run `python -m pip install xgboost` to enable it."
        )

    return pd.concat(validation_predictions, ignore_index=True), pd.DataFrame(selected)


def retrain_and_test(train_validation: pd.DataFrame, test: pd.DataFrame, selected: pd.DataFrame):
    outputs = evaluate_baselines(test)

    for row in selected.to_dict("records"):
        structure = row["structure"]
        model_name = row["model"]
        display_model = row["display_model"]
        params = json.loads(row["parameters"])
        started = time.perf_counter()

        if structure == "Local":
            fit_load, test_load, _ = fit_local_predict(model_name, params, train_validation, test)
        else:
            fit_load, test_load, _ = fit_global_predict(model_name, params, train_validation, test)

        elapsed = time.perf_counter() - started
        fit_pred = to_demand_prediction(fit_load, train_validation)
        test_pred = to_demand_prediction(test_load, test)
        fit_metrics = metric_dict(train_validation["demand"], fit_pred)
        test_metrics = metric_dict(test["demand"], test_pred)
        add_training_log("final_retrain_test", structure, display_model, params, elapsed, fit_metrics, test_metrics)
        log(
            f"FINAL {structure}-{display_model} params={params} | time={elapsed:.3f}s | "
            f"train+validation WAPE={fit_metrics['WAPE']:.4f} | test WAPE={test_metrics['WAPE']:.4f}"
        )
        outputs.append(prediction_frame(test, structure, display_model, test_pred))

    return pd.concat(outputs, ignore_index=True)


def format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    sub_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9E2F3")

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.border = Border(bottom=thin)
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
        for column_cells in sheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells[:200]]
            width = min(max(max(map(len, values), default=8) + 2, 10), 28)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

        if sheet.title == "Summary":
            for row in sheet.iter_rows(min_row=2):
                row[0].fill = sub_fill
                row[1].fill = sub_fill
        if sheet.title == "Predictions":
            for cell in sheet["C"][1:]:
                cell.number_format = "yyyy-mm"

    workbook.save(path)


def write_results_workbook(
    path: Path,
    predictions: pd.DataFrame,
    selected: pd.DataFrame,
    split_description: pd.DataFrame,
    logs: pd.DataFrame,
) -> None:
    summary = summary_metrics(predictions)
    dept = department_metrics(predictions)
    predictions = predictions.sort_values(["structure", "model", "date", "department"])

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        dept.to_excel(writer, sheet_name="Department_Metrics", index=False)
        predictions.to_excel(writer, sheet_name="Predictions", index=False)
        selected.to_excel(writer, sheet_name="Selected_Params", index=False)
        split_description.to_excel(writer, sheet_name="Data_Split", index=False)
        logs.to_excel(writer, sheet_name="Training_Log", index=False)
    format_workbook(path)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    log(f"Reading data: {DATA_PATH}")
    raw = load_excel_data(DATA_PATH)
    featured = add_features(raw)
    train, validation, test = split_data(featured)
    log(
        f"Samples: train={len(train)}, validation={len(validation)}, test={len(test)}; "
        f"departments={raw['department'].nunique()}"
    )

    validation_predictions, selected = tune_and_validate(train, validation)
    validation_logs = pd.DataFrame(TRAINING_LOGS)

    split_description = pd.DataFrame(
        [
            {"set": "Training", "target_months": "2023-04 to 2024-12", "samples": len(train)},
            {"set": "Validation", "target_months": "2025-01 to 2025-06", "samples": len(validation)},
            {"set": "Test", "target_months": "2025-07 to 2025-12", "samples": len(test)},
        ]
    )
    write_results_workbook(
        VALIDATION_XLSX,
        validation_predictions,
        selected,
        split_description,
        validation_logs,
    )
    log(f"Validation workbook written: {VALIDATION_XLSX}")

    train_validation = pd.concat([train, validation], ignore_index=True)
    test_predictions = retrain_and_test(train_validation, test, selected)
    all_logs = pd.DataFrame(TRAINING_LOGS)
    write_results_workbook(
        TEST_XLSX,
        test_predictions,
        selected,
        split_description,
        all_logs,
    )
    log(f"Test workbook written: {TEST_XLSX}")
    log("Completed successfully.")


if __name__ == "__main__":
    main()
