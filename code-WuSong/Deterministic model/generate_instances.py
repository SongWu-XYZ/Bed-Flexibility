"""由“算例生成输入数据.xlsx”生成81组参数×12个月的30科室算例。"""

from __future__ import annotations

import csv
import json
from calendar import monthrange
from itertools import product
from pathlib import Path

from openpyxl import load_workbook


# ======================= 用户设置 =======================
ROOT = Path(__file__).resolve().parent
INPUT_FILE = ROOT / "算例生成输入数据.xlsx"
SHEET_NAME = "算例生成输入数据"
OUTPUT_DIR = ROOT / "generated_instances"
DATA_YEAR = 2023
BASE_DAILY_WAGE = 200.0
NURSE_BED_RATIO = 0.5
WAGE_FACTORS = (0.90, 1.00, 1.10)
MULTISKILL_COST_RATIOS = (0.10, 0.20, 0.30)
ASSIGNMENT_COST_RATIOS = (0.05, 0.10, 0.15)
DEMAND_FACTORS = (0.90, 1.00, 1.10)
# ========================================================

MONTHS = range(1, 13)
MONTH_NAMES = ("一月", "二月", "三月", "四月", "五月", "六月",
               "七月", "八月", "九月", "十月", "十一月", "十二月")
MONTH_FIELDS = ("入院人数", "服务能力", "LOS'", "需求", "单日拒绝成本")
DEPARTMENTS = (
    "心血管内科", "血液内科", "神经内科", "呼吸内科", "消化内科", "内分泌科",
    "肾病内科", "康复医学科", "风湿免疫科", "运动医学与小儿骨科",
    "脊柱与骨肿瘤科", "创伤与手外科", "神经外科", "普通外科", "小儿外科",
    "肿瘤科", "胸外科", "心外科", "泌尿外科", "放疗科", "骨关节外科",
    "移植外科", "整形美容烧伤", "感染科", "小儿内科", "眼科", "耳鼻喉科",
    "皮肤科", "中医科", "全科医学科",
)


def number(value, location: str) -> float:
    if value is None:
        raise ValueError(f"{location}为空；若为公式，请先用Excel打开并保存工作簿")
    try:
        ans = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{location}不是数值：{value!r}") from exc
    if ans < 0:
        raise ValueError(f"{location}不能为负数：{ans}")
    return ans


def read_departments() -> list[dict]:
    if not INPUT_FILE.is_file():
        raise FileNotFoundError(f"找不到输入文件：{INPUT_FILE}")
    book = load_workbook(INPUT_FILE, read_only=True, data_only=True)
    if SHEET_NAME not in book.sheetnames:
        raise ValueError(f"找不到工作表{SHEET_NAME!r}：{book.sheetnames}")
    sheet = book[SHEET_NAME]
    if tuple(sheet.cell(1, c).value for c in range(1, 4)) != ("科室", "床位数", "拒绝成本"):
        raise ValueError("前三列必须依次为：科室、床位数、拒绝成本")
    for m, month_name in enumerate(MONTH_NAMES, 1):
        start = 4 + 5 * (m - 1)
        if sheet.cell(1, start).value != month_name:
            raise ValueError(f"第{start}列应为{month_name}")
        fields = tuple(sheet.cell(2, start + k).value for k in range(5))
        if fields != MONTH_FIELDS:
            raise ValueError(f"{month_name}字段错误：{fields}")

    result = []
    for dept_id, expected_name in enumerate(DEPARTMENTS, 1):
        row = dept_id + 2
        dept_name = str(sheet.cell(row, 1).value or "").strip()
        if dept_name != expected_name:
            raise ValueError(f"第{row}行应为{expected_name}，实际为{dept_name}")
        beds_raw = number(sheet.cell(row, 2).value, f"{dept_name}-床位数")
        if abs(beds_raw - round(beds_raw)) > 1e-9:
            raise ValueError(f"{dept_name}床位数必须为整数")
        beds = int(round(beds_raw))
        rejection = number(sheet.cell(row, 3).value, f"{dept_name}-拒绝成本")
        monthly = []
        for month in MONTHS:
            start = 4 + 5 * (month - 1)
            tag = f"{dept_name}-{MONTH_NAMES[month - 1]}"
            patients_raw = number(sheet.cell(row, start).value, tag + "-入院人数")
            if abs(patients_raw - round(patients_raw)) > 1e-9:
                raise ValueError(f"{tag}-入院人数必须为整数")
            patients = int(round(patients_raw))
            capacity = number(sheet.cell(row, start + 1).value, tag + "-服务能力")
            los = number(sheet.cell(row, start + 2).value, tag + "-LOS'")
            demand = number(sheet.cell(row, start + 3).value, tag + "-需求")
            daily_rejection = number(sheet.cell(row, start + 4).value, tag + "-单日拒绝成本")
            days = monthrange(DATA_YEAR, month)[1]
            if abs(capacity - beds * days) > 1e-6:
                raise ValueError(f"{tag}-服务能力应为{beds * days}，实际为{capacity}")
            if los <= 0 or abs(demand - patients * los) > 1e-5:
                raise ValueError(f"{tag}-需求不等于入院人数×LOS'")
            if abs(daily_rejection - rejection / los) > 1e-5:
                raise ValueError(f"{tag}-单日拒绝成本不等于拒绝成本/LOS'")
            monthly.append({"patients": patients, "capacity": capacity, "los": los,
                            "demand": demand, "daily_rejection": daily_rejection})
        result.append({"id": dept_id, "beds": beds, "rejection": rejection,
                       "monthly": monthly})
    book.close()
    return result


def parameter_id(wf, qr, ar, df) -> str:
    return (f"W{round(wf * 100):03d}_Q{round(qr * 100):02d}_"
            f"A{round(ar * 100):02d}_D{round(df * 100):03d}")


def write_txt(path: Path, instance: dict) -> None:
    p = instance["parameters"]
    with path.open("w", encoding="utf-8", newline="\n") as out:
        out.write(f"scenario_id {instance['scenario_id']}\n")
        out.write(f"month {instance['month']}\n")
        out.write(f"departments {len(instance['departments'])}\n")
        out.write(f"nurse_bed_ratio {p['nurse_bed_ratio']:.2f}\n")
        out.write(f"days_per_month {p['days_per_month']}\n")
        out.write(f"daily_wage {p['daily_wage']:.2f}\n")
        out.write(f"multiskill_cost_q {p['multiskill_cost_q']:.2f}\n")
        out.write(f"assignment_cost_ratio {p['assignment_cost_ratio']:.2f}\n")
        out.write(f"demand_factor {p['demand_factor']:.2f}\n")
        out.write("# id|beds|rejection_cost|daily_rejection_cost|daily_assignment_cost|"
                  "nominal_los|patient_demand|demand_bed_days|capacity_bed_days\n")
        for d in instance["departments"]:
            out.write(f"{d['department_id']}|{d['beds']}|{d['rejection_cost']:.2f}|"
                      f"{d['daily_rejection_cost']:.2f}|{d['daily_assignment_cost']:.2f}|"
                      f"{d['nominal_los']:.2f}|{d['patient_demand']}|"
                      f"{d['demand']:.2f}|{d['capacity']:.2f}\n")


def main() -> None:
    source = read_departments()
    json_dir, txt_dir = OUTPUT_DIR / "json", OUTPUT_DIR / "txt"
    json_dir.mkdir(parents=True, exist_ok=True)
    txt_dir.mkdir(parents=True, exist_ok=True)
    for folder, pattern in ((json_dir, "*.json"), (txt_dir, "*.txt")):
        for old_file in folder.glob(pattern):
            old_file.unlink()

    index_rows, long_rows = [], []
    grid = product(WAGE_FACTORS, MULTISKILL_COST_RATIOS,
                   ASSIGNMENT_COST_RATIOS, DEMAND_FACTORS)
    for wf, qr, ar, df in grid:
        pid = parameter_id(wf, qr, ar, df)
        wage, q = BASE_DAILY_WAGE * wf, BASE_DAILY_WAGE * wf * qr
        for month in MONTHS:
            days, sid = monthrange(DATA_YEAR, month)[1], f"{pid}_M{month:02d}"
            instance_departments = []
            for d in source:
                m = d["monthly"][month - 1]
                item = {
                    "department_id": d["id"], "beds": d["beds"],
                    "rejection_cost": d["rejection"],
                    "daily_rejection_cost": m["daily_rejection"],
                    "daily_assignment_cost": m["daily_rejection"] * ar,
                    "nominal_los": m["los"], "base_patient_demand": m["patients"],
                    "patient_demand": max(0, round(m["patients"] * df)),
                    "base_demand": m["demand"], "demand": m["demand"] * df,
                    "capacity": m["capacity"],
                }
                instance_departments.append(item)
                long_rows.append({"scenario_id": sid, "parameter_id": pid,
                    "year": DATA_YEAR, "month": month, "days_per_month": days,
                    "department_id": d["id"], "beds": d["beds"],
                    "rejection_cost": d["rejection"],
                    "daily_rejection_cost": item["daily_rejection_cost"],
                    "daily_assignment_cost": item["daily_assignment_cost"],
                    "nominal_los": item["nominal_los"],
                    "base_patient_demand": item["base_patient_demand"],
                    "patient_demand": item["patient_demand"],
                    "base_demand_bed_days": item["base_demand"],
                    "demand_bed_days": item["demand"],
                    "capacity_bed_days": item["capacity"],
                    "daily_wage": wage, "multiskill_cost_q": q})
            parameters = {"data_year": DATA_YEAR, "days_per_month": days,
                "nurse_bed_ratio": NURSE_BED_RATIO, "base_daily_wage": BASE_DAILY_WAGE,
                "wage_factor": wf, "daily_wage": wage,
                "multiskill_cost_ratio": qr, "multiskill_cost_q": q,
                "assignment_cost_ratio": ar, "demand_factor": df}
            instance = {"scenario_id": sid, "parameter_id": pid, "year": DATA_YEAR,
                        "month": month, "parameters": parameters,
                        "departments": instance_departments}
            (json_dir / f"{sid}.json").write_text(
                json.dumps(instance, ensure_ascii=False, indent=2), encoding="utf-8")
            write_txt(txt_dir / f"{sid}.txt", instance)
            index_rows.append({"scenario_id": sid, "parameter_id": pid,
                               "year": DATA_YEAR, "month": month, **parameters})

    for path, rows in ((OUTPUT_DIR / "scenario_index.csv", index_rows),
                       (OUTPUT_DIR / "instances_long.csv", long_rows)):
        with path.open("w", encoding="utf-8-sig", newline="") as out:
            writer = csv.DictWriter(out, fieldnames=rows[0].keys())
            writer.writeheader(); writer.writerows(rows)
    print(f"生成完成：{len(index_rows)}个算例 = 81组参数×12个月；每例30个科室。")
    print(f"年份={DATA_YEAR}；使用各月实际天数；输出目录：{OUTPUT_DIR}")


if __name__ == "__main__":
    main()
